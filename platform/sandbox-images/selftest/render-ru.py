#!/usr/bin/env python3
"""Render Cyrillic documents, inspect embedded fonts, and OCR page images."""
import argparse
from collections import Counter
import json
from pathlib import Path
import re
import resource
import shutil
import subprocess
import tempfile
import time
import xml.etree.ElementTree as ET


def command(args):
    return subprocess.check_output([str(arg) for arg in args], text=True, stderr=subprocess.PIPE, timeout=120)


def words(text):
    return Counter(re.findall(r'[а-яё]+', text.lower().replace('ё', 'е')))


def recall(expected, actual):
    return sum((words(expected) & words(actual)).values()) / sum(words(expected).values())


def inspect_pdf(pdf, expected, negative=False):
    prefix = pdf.with_suffix('')
    command(['pdftoppm', '-r', '150', '-png', pdf, prefix])
    pages = sorted(pdf.parent.glob(prefix.name + '-*.png'))
    assert pages, 'No raster pages'
    text = '\n'.join(command(['tesseract', page, 'stdout', '-l', 'rus']) for page in pages)
    score = recall(expected, text)
    if negative:
        assert score < 0.5, f'Negative control unexpectedly readable: {score}'
        return {'word_recall': score, 'passed': True}
    assert score >= 0.95, f'{pdf.name}: OCR recall {score:.3f}, expected >= 0.95'
    font_output = command(['pdffonts', pdf])
    assert 'Type 3' not in font_output, 'Type 3 font found'
    font_lines = font_output.splitlines()[2:]
    assert font_lines, 'No fonts found'
    assert all(line.split()[-5] == 'yes' for line in font_lines), 'Font is not embedded'
    assert all(re.search(r'Carlito|Liberation|DejaVu', line) for line in font_lines), font_output
    layout = command(['pdftotext', '-layout', pdf, '-'])
    # The heading and pangram alone take two lines. The long paragraph must wrap.
    paragraph_words = set(words(expected.splitlines()[-1]))
    paragraph_lines = [line for line in layout.splitlines() if set(words(line)) & paragraph_words]
    assert len(paragraph_lines) > 1, 'Paragraph did not wrap'
    bbox = ET.fromstring(command(['pdftotext', '-bbox', pdf, '-']))
    ns = {'p': 'http://www.w3.org/1999/xhtml'}
    for page in bbox.findall('.//p:page', ns):
        width, height = float(page.attrib['width']), float(page.attrib['height'])
        for word in page.findall('.//p:word', ns):
            assert 0 <= float(word.attrib['xMin']) < float(word.attrib['xMax']) <= width, 'Text outside page width'
            assert 0 <= float(word.attrib['yMin']) < float(word.attrib['yMax']) <= height, 'Text outside page height'
    return {'word_recall': score, 'embedded_fonts': font_lines, 'paragraph_lines': len(paragraph_lines), 'passed': True}


def make_documents(directory, lines):
    from docx import Document
    from docx.shared import Pt
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from pptx import Presentation
    from pptx.util import Inches, Pt as SlidePt

    document = Document()
    for text, font in zip(lines, ['Calibri', 'Times New Roman', 'Calibri']):
        run = document.add_paragraph().add_run(text)
        run.font.name, run.font.size = font, Pt(16)
    document.save(directory / 'sample.docx')

    book = Workbook()
    sheet = book.active
    for row, text in enumerate(lines, 1):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        cell = sheet.cell(row, 1, text)
        cell.font = Font(name='Arial', size=16)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        sheet.row_dimensions[row].height = [40, 90, 240][row-1]
    sheet.print_area = 'A1:H3'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 1
    book.save(directory / 'sample.xlsx')

    slides = Presentation()
    frame = slides.slides.add_slide(slides.slide_layouts[6]).shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8.8), Inches(6.3)).text_frame
    frame.word_wrap = True
    for i, text in enumerate(lines):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.text = text
        paragraph.font.name, paragraph.font.size = 'Calibri', SlidePt(20)
    slides.save(directory / 'sample.pptx')


def workloads(directory):
    from pptx import Presentation
    import pandas as pd

    samples = {}
    start = time.monotonic()
    slides = Presentation()
    for i in range(20):
        slide = slides.slides.add_slide(slides.slide_layouts[1])
        slide.shapes.title.text = f'Страница {i+1}'
        slide.placeholders[1].text = 'Проверка памяти при создании презентации'
    slides.save(directory / 'twenty-slides.pptx')
    samples['twenty_slides_seconds'] = time.monotonic() - start
    start = time.monotonic()
    table = pd.DataFrame({'number': range(100000), 'group': [i % 10 for i in range(100000)]})
    assert table.groupby('group')['number'].sum().sum() == 4999950000
    table.to_csv(directory / 'data.csv', index=False)
    samples['pandas_100000_rows_seconds'] = time.monotonic() - start
    return samples


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--json', action='store_true')
    parser.add_argument('--workloads', action='store_true')
    args = parser.parse_args()
    lines = Path(__file__).with_name('sample-ru.txt').read_text().strip().splitlines()
    expected = '\n'.join(lines)
    report = {'formats': {}, 'notDone': []}
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from reportlab.pdfgen import canvas

    # All generated files stay inside the disposable container or VM.
    with tempfile.TemporaryDirectory(prefix='oma-render-') as name:
        directory = Path(name)
        pdfmetrics.registerFont(TTFont('DejaVuSans', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
        style = ParagraphStyle('ru', fontName='DejaVuSans', fontSize=16, leading=24)
        story = []
        for line in lines:
            story += [Paragraph(line, style), Spacer(1, 12)]
        pdf = directory / 'reportlab.pdf'
        # The document template passes its own initial font to the canvas and
        # overrides a canvasmaker argument, so Helvetica would stay in the PDF.
        SimpleDocTemplate(str(pdf), initialFontName='DejaVuSans').build(story)
        report['formats']['pdf'] = inspect_pdf(pdf, expected)
        control = directory / 'negative.pdf'
        drawing = canvas.Canvas(str(control))
        drawing.setFont('Helvetica', 16)
        for i, line in enumerate(lines):
            drawing.drawString(40, 750-i*50, line)
        drawing.save()
        report['negative_control'] = inspect_pdf(control, expected, negative=True)
        if shutil.which('soffice'):
            make_documents(directory, lines)
            for ext in ['docx', 'xlsx', 'pptx']:
                out = directory / ext
                out.mkdir()
                start = time.monotonic()
                command(['soffice', '-env:UserInstallation=file://' + str(directory / 'lo-profile'), '--headless', '--convert-to', 'pdf', '--outdir', out, directory / ('sample.' + ext)])
                converted = out / 'sample.pdf'
                assert converted.exists(), f'LibreOffice produced no PDF for {ext}'
                report['formats'][ext] = inspect_pdf(converted, expected)
                report['formats'][ext]['convert_and_check_seconds'] = time.monotonic() - start
        else:
            report['notDone'] = ['docx, xlsx and pptx rendering require the office target']
        if args.workloads:
            report['workloads'] = workloads(directory)
        report['python_peak_rss_kib'] = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
        report['child_peak_rss_kib'] = resource.getrusage(resource.RUSAGE_CHILDREN).ru_maxrss
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
